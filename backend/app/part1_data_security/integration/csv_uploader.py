import csv
import datetime
import re
from sqlalchemy import text
from app.db.database import engine

# ============================================
# ข้อจำกัดด้านความปลอดภัย
# ============================================
MAX_FILE_SIZE_MB = 10
ALLOWED_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xls"}

# รายการ Encoding ที่รองรับ (เรียงตามลำดับความน่าจะเป็น)
# cp874 และ tis-620 เป็น Encoding ภาษาไทยที่ Microsoft Excel บน Windows ใช้เป็นค่าเริ่มต้น
ENCODINGS_TO_TRY = ["utf-8-sig", "cp874", "tis-620", "utf-8", "windows-1252", "latin-1"]

# SQL Reserved Words ที่ห้ามใช้เป็นชื่อตาราง/คอลัมน์
_SQL_RESERVED = {
    "select", "insert", "update", "delete", "drop", "create", "alter",
    "table", "from", "where", "and", "or", "not", "null", "index",
    "pragma", "attach", "detach", "vacuum", "grant", "revoke",
    "order", "group", "by", "limit", "join", "having", "union",
}


def sanitize_identifier(name: str) -> str:
    """
    Sanitize SQL identifier (table name / column name) to prevent SQL Injection.
    - รองรับทั้งตัวอักษรภาษาไทย (\u0E00-\u0E7F) และภาษาอังกฤษ (a-z, A-Z), ตัวเลข (0-9) และ underscore
    - ลบอักขระพิเศษอันตราย เช่น ;, ', ", --, /*, (), =, <, > ออกทั้งหมด
    - ถ้าขึ้นต้นด้วยตัวเลข จะเติม 'col_' นำหน้า
    - ถ้าตรงกับ SQL Reserved Word จะเติม 'col_' นำหน้า
    - ถ้าว่างเปล่า จะใช้ชื่อ 'unnamed'
    """
    clean = re.sub(r'[^\w\u0E00-\u0E7F]', '_', name.strip()).lower()
    # ลบ underscore ซ้ำ
    clean = re.sub(r'_+', '_', clean).strip('_')
    if not clean:
        clean = "unnamed"
    if clean[0].isdigit():
        clean = f"col_{clean}"
    if clean in _SQL_RESERVED:
        clean = f"col_{clean}"
    return clean


def _read_csv_with_fallback(file_path: str):
    """
    อ่านไฟล์ CSV โดยพยายามใช้ Encoding ต่างๆ (UTF-8, CP874 สำหรับภาษาไทยจาก Excel, TIS-620, Windows-1252)
    ป้องกัน UnicodeDecodeError เมื่อผู้ใช้อัปโหลดไฟล์ที่บันทึกจาก Excel บน Windows
    """
    last_error = None
    for enc in ENCODINGS_TO_TRY:
        try:
            with open(file_path, mode="r", encoding=enc, newline="") as f:
                reader = csv.reader(f)
                try:
                    raw_headers = next(reader)
                except StopIteration:
                    return [], [], enc
                rows = list(reader)
                return raw_headers, rows, enc
        except (UnicodeDecodeError, UnicodeError) as e:
            last_error = e
            continue
        except Exception as e:
            last_error = e
            break

    raise ValueError(f"ไม่สามารถถอดรหัสตัวอักษรของไฟล์ CSV ได้ (สาเหตุ: {last_error}) กรุณาบันทึกเป็น UTF-8 หรือ Windows Thai (CP874)")


def _read_tsv_txt_with_fallback(file_path: str, ext: str):
    """
    อ่านไฟล์ TSV หรือ Text ที่คั่นด้วย Tab, Pipe, Comma หรือ Semicolon
    พร้อม Auto-Encoding Detection สำหรับภาษาไทย
    """
    last_error = None
    for enc in ENCODINGS_TO_TRY:
        try:
            with open(file_path, mode="r", encoding=enc, newline="") as f:
                if ext == ".tsv":
                    delimiter = "\t"
                else:
                    sample = f.read(2048)
                    f.seek(0)
                    if "\t" in sample:
                        delimiter = "\t"
                    elif "|" in sample:
                        delimiter = "|"
                    elif ";" in sample:
                        delimiter = ";"
                    else:
                        delimiter = ","

                reader = csv.reader(f, delimiter=delimiter)
                try:
                    raw_headers = next(reader)
                except StopIteration:
                    return [], [], enc
                rows = list(reader)
                return raw_headers, rows, enc
        except (UnicodeDecodeError, UnicodeError) as e:
            last_error = e
            continue
        except Exception as e:
            last_error = e
            break

    raise ValueError(f"ไม่สามารถถอดรหัสตัวอักษรของไฟล์ {ext} ได้ (สาเหตุ: {last_error})")


def _read_excel(file_path: str, ext: str):
    """
    อ่านไฟล์ Excel (.xlsx, .xls) โดยดึงข้อมูลจากชีตแรก (Active Sheet)
    รองรับภาษาไทย และแปลงวันที่และตัวเลขให้อยู่ในรูปแบบที่ถูกต้อง
    """
    if ext == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise ImportError("ไม่พบไลบรารี openpyxl สำหรับอ่านไฟล์ .xlsx (กรุณาติดตั้งด้วย pip install openpyxl)")

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            ws = wb.active
            raw_headers = []
            rows = []

            for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if r_idx == 0:
                    raw_headers = [
                        str(cell).strip() if cell is not None and str(cell).strip() != ""
                        else f"col_{i+1}"
                        for i, cell in enumerate(row)
                    ]
                else:
                    if all(cell is None or str(cell).strip() == "" for cell in row):
                        continue
                    formatted_row = []
                    for cell in row:
                        if cell is None:
                            formatted_row.append("")
                        elif isinstance(cell, (datetime.datetime, datetime.date)):
                            formatted_row.append(cell.strftime("%Y-%m-%d"))
                        elif isinstance(cell, float):
                            if cell.is_integer():
                                formatted_row.append(str(int(cell)))
                            else:
                                formatted_row.append(str(cell))
                        else:
                            formatted_row.append(str(cell).strip())

                    if len(formatted_row) < len(raw_headers):
                        formatted_row += [""] * (len(raw_headers) - len(formatted_row))
                    rows.append(formatted_row[:len(raw_headers)])

            wb.close()
            return raw_headers, rows, "Excel (.xlsx) - UTF-8"
        except Exception as e:
            raise ValueError(f"เกิดข้อผิดพลาดในการอ่านไฟล์ Excel (.xlsx): {str(e)}")

    elif ext == ".xls":
        try:
            import xlrd
            wb = xlrd.open_workbook(file_path)
            sheet = wb.sheet_by_index(0)
            if sheet.nrows == 0:
                return [], [], "Excel (.xls)"
            raw_headers = [str(sheet.cell_value(0, c)).strip() or f"col_{c+1}" for c in range(sheet.ncols)]
            rows = []
            for r in range(1, sheet.nrows):
                row_vals = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                if all(v == "" for v in row_vals):
                    continue
                rows.append(row_vals)
            return raw_headers, rows, "Excel (.xls)"
        except ImportError:
            raise ValueError("ระบบรองรับไฟล์ Excel รุ่นใหม่ (.xlsx) เป็นหลัก กรุณาบันทึกไฟล์เป็น .xlsx หรือติดตั้ง xlrd")
        except Exception as e:
            raise ValueError(f"เกิดข้อผิดพลาดในการอ่านไฟล์ .xls: {str(e)}")

    raise ValueError(f"นามสกุลไฟล์ {ext} ไม่รองรับสำหรับการอ่านแบบ Excel")


def _read_file_data(file_path: str):
    """Dispatcher เพื่ออ่านข้อมูลตามนามสกุลไฟล์"""
    ext = "." + file_path.rsplit(".", 1)[-1].lower() if "." in file_path else ""
    if ext == ".csv":
        return _read_csv_with_fallback(file_path)
    elif ext in (".tsv", ".txt"):
        return _read_tsv_txt_with_fallback(file_path, ext)
    elif ext in (".xlsx", ".xls"):
        return _read_excel(file_path, ext)
    else:
        raise ValueError(f"นามสกุลไฟล์ {ext} ไม่รองรับ (อนุญาตเฉพาะ {', '.join(ALLOWED_EXTENSIONS)})")


# ชุดค่าว่างและ placeholder ที่พบบ่อยในไฟล์ CSV เช่น -, N/A, NA, null, nil
NA_PLACEHOLDERS = {"", "-", "--", "—", "n/a", "na", "null", "none", "nil", "nan", "."}



def _detect_column_type(values: list) -> str:
    """
    ตรวจจับชนิดข้อมูลอัตโนมัติจากค่าตัวอย่างในคอลัมน์
    ลำดับการตรวจ: INTEGER -> REAL -> DATE -> TEXT
    - กรองค่าว่างและ NA Placeholders (เช่น "-", "N/A", "null") ออกก่อนตรวจจับ
    - หากค่าตัวอย่างที่ถูกต้องทั้งหมดเป็นตัวเลข จะระบุเป็น INTEGER หรือ REAL ได้อย่างแม่นยำ
    """
    # กรองค่าว่างและ NA placeholders ออก
    samples = [
        v.strip() for v in values 
        if v and v.strip() and v.strip().lower() not in NA_PLACEHOLDERS
    ]
    if not samples:
        return "TEXT"

    # ตรวจสอบ INTEGER
    is_int = True
    for s in samples:
        try:
            int(s.replace(",", ""))  # รองรับตัวเลขที่มีจุลภาค เช่น 1,000
        except ValueError:
            is_int = False
            break
    if is_int:
        return "INTEGER"

    # ตรวจสอบ REAL (float)
    is_float = True
    for s in samples:
        try:
            float(s.replace(",", ""))
        except ValueError:
            is_float = False
            break
    if is_float:
        return "REAL"

    # ตรวจสอบ DATE (รูปแบบ YYYY-MM-DD หรือ DD/MM/YYYY หรือ YYYY/MM/DD)
    date_pattern = re.compile(
        r"^\d{4}-\d{2}-\d{2}$|^\d{2}/\d{2}/\d{4}$|^\d{4}/\d{2}/\d{2}$"
    )
    is_date = all(date_pattern.match(s) for s in samples)
    if is_date:
        return "DATE"

    return "TEXT"


def _normalize_date(date_str: str) -> str:
    """แปลงรูปแบบวันที่ DD/MM/YYYY หรือ YYYY/MM/DD เป็น ISO YYYY-MM-DD สำหรับ SQLite"""
    if not date_str or not isinstance(date_str, str):
        return date_str
    date_str = date_str.strip()
    # ถ้าเป็น YYYY-MM-DD อยู่แล้ว ไม่ต้องแปลง
    if re.match(r'^\d{4}-\d{2}-\d{2}$', date_str):
        return date_str
    # DD/MM/YYYY → YYYY-MM-DD
    m = re.match(r'^(\d{2})/(\d{2})/(\d{4})$', date_str)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    # YYYY/MM/DD → YYYY-MM-DD
    m = re.match(r'^(\d{4})/(\d{2})/(\d{2})$', date_str)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return date_str


def validate_file_extension(filename: str) -> bool:
    """ตรวจสอบนามสกุลไฟล์ว่าอนุญาตหรือไม่"""
    if not filename:
        return False
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in ALLOWED_EXTENSIONS


def validate_file_size(file_path: str) -> bool:
    """ตรวจสอบขนาดไฟล์ว่าไม่เกินขีดจำกัด"""
    import os
    size_mb = os.path.getsize(file_path) / (1024 * 1024)
    return size_mb <= MAX_FILE_SIZE_MB


def detect_pii(headers: list, rows: list) -> list:
    """
    ตรวจจับข้อมูลส่วนบุคคลที่ละเอียดอ่อน (Personally Identifiable Information: PII)
    เช่น เลขบัตรประจำตัวประชาชน, เบอร์โทรศัพท์, หรืออีเมล
    คืนค่ารายการประเภท PII และคอลัมน์ที่ตรวจพบเพื่อแจ้งเตือนผู้ใช้งาน
    """
    warnings = []
    id_card_regex = re.compile(r'^\d{13}$|^\d{1}-\d{4}-\d{5}-\d{2}-\d{1}$')
    phone_regex = re.compile(r'^0[2689]\d{7,8}$|^0[2689]-\d{3,4}-\d{4}$')
    email_regex = re.compile(r'^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$')

    for col_idx, col_name in enumerate(headers):
        col_samples = [
            row[col_idx].strip()
            for row in rows[:50]
            if col_idx < len(row) and row[col_idx] and row[col_idx].strip()
        ]
        if not col_samples:
            continue

        # 1. ตรวจสอบเลขประจำตัวประชาชน (13 หลัก)
        if any(id_card_regex.match(s) for s in col_samples):
            warnings.append({
                "column": col_name,
                "type": "citizen_id",
                "message": f"คอลัมน์ '{col_name}' มีข้อมูลที่มีลักษณะคล้ายเลขบัตรประจำตัวประชาชน (13 หลัก)"
            })
            continue

        # 2. ตรวจสอบเบอร์โทรศัพท์
        if any(phone_regex.match(s) for s in col_samples):
            warnings.append({
                "column": col_name,
                "type": "phone_number",
                "message": f"คอลัมน์ '{col_name}' มีข้อมูลที่มีลักษณะคล้ายเบอร์โทรศัพท์"
            })
            continue

        # 3. ตรวจสอบอีเมล
        if any(email_regex.match(s) for s in col_samples):
            warnings.append({
                "column": col_name,
                "type": "email",
                "message": f"คอลัมน์ '{col_name}' มีข้อมูลที่อยู่อีเมล"
            })
            continue

    return warnings


def upload_file_to_db(file_path: str, table_name: str) -> dict:
    """
    ระบบการนำเข้าข้อมูลและจัดการโครงสร้าง (Data Integration & Schema Mapping System)
    - รองรับไฟล์ CSV, TSV, TXT, Excel (.xlsx, .xls)
    - Auto-Encoding Detection สำหรับภาษาไทย (UTF-8, CP874, TIS-620)
    - ตรวจจับชนิดข้อมูลอัตโนมัติ (INTEGER, REAL, DATE, TEXT)
    - ป้องกัน SQL Injection ด้วยการ Sanitize ชื่อตารางและชื่อคอลัมน์ (รองรับภาษาไทย)
    - ตรวจสอบนามสกุลไฟล์และขนาดไฟล์
    - นำเข้าตารางในฐานข้อมูล SQLite
    """
    try:
        # --- ตรวจสอบนามสกุลไฟล์ ---
        if not validate_file_extension(file_path):
            return {
                "status": "error",
                "message": f"File type not allowed. Accepted: {', '.join(ALLOWED_EXTENSIONS)}"
            }

        # --- ตรวจสอบขนาดไฟล์ ---
        if not validate_file_size(file_path):
            return {
                "status": "error",
                "message": f"File too large. Max size: {MAX_FILE_SIZE_MB} MB"
            }

        table_clean = sanitize_identifier(table_name)

        # --- อ่านข้อมูลไฟล์ตามนามสกุลพร้อม Fallback อัตโนมัติ ---
        try:
            raw_headers, rows, used_encoding = _read_file_data(file_path)
        except Exception as e:
            return {
                "status": "error",
                "message": str(e)
            }

        if not raw_headers:
            return {"status": "error", "message": "ไฟล์ที่อัปโหลดไม่มีหัวตาราง (Headers)"}

        if not rows:
            return {"status": "error", "message": "ไฟล์ที่อัปโหลดไม่มีข้อมูล (Data rows)"}

        # --- Sanitize ชื่อคอลัมน์เพื่อป้องกัน SQL Injection ---
        headers = []
        seen = {}
        for h in raw_headers:
            clean = sanitize_identifier(h)
            # จัดการชื่อคอลัมน์ซ้ำ: เติมเลขต่อท้าย
            if clean in seen:
                seen[clean] += 1
                clean = f"{clean}_{seen[clean]}"
            else:
                seen[clean] = 0
            headers.append(clean)

        # --- ตรวจจับชนิดข้อมูลอัตโนมัติ (Data Type Detection) ---
        col_types = {}
        for col_idx, col_name in enumerate(headers):
            col_values = [row[col_idx] for row in rows if col_idx < len(row)]
            col_types[col_name] = _detect_column_type(col_values)

        col_defs = ", ".join([f'"{col}" {col_types[col]}' for col in headers])
        drop_sql = f'DROP TABLE IF EXISTS "{table_clean}";'
        create_sql = f'CREATE TABLE "{table_clean}" ({col_defs});'

        quoted_headers = [f'"{col}"' for col in headers]
        placeholders = ", ".join([f":param_{i}" for i in range(len(headers))])
        insert_sql = f'INSERT INTO "{table_clean}" ({", ".join(quoted_headers)}) VALUES ({placeholders});'

        # --- นำเข้าฐานข้อมูล ---
        inserted_count = 0
        with engine.connect() as conn:
            conn.execute(text(drop_sql))
            conn.execute(text(create_sql))
            batch_rows = []
            for row in rows:
                if len(row) == len(headers):
                    row_dict = {}
                    for i, col in enumerate(headers):
                        val = row[i].strip() if row[i] else None
                        # แปลงค่าตามชนิดที่ตรวจจับได้
                        if val is not None and val != "":
                            clean_str = val.strip()
                            if clean_str.lower() in NA_PLACEHOLDERS:
                                val = None
                            else:
                                clean_num = clean_str.replace(",", "")
                                if col_types[col] == "INTEGER":
                                    try:
                                        val = int(clean_num)
                                    except ValueError:
                                        try:
                                            val = int(float(clean_num))
                                        except ValueError:
                                            pass
                                elif col_types[col] == "REAL":
                                    try:
                                        val = float(clean_num)
                                    except ValueError:
                                        pass
                                elif col_types[col] == "DATE":
                                    # Normalize วันที่เป็น ISO format YYYY-MM-DD สำหรับ SQLite
                                    val = _normalize_date(clean_str)
                        else:
                            val = None
                        row_dict[f"param_{i}"] = val
                    batch_rows.append(row_dict)

            if batch_rows:
                conn.execute(text(insert_sql), batch_rows)
                inserted_count = len(batch_rows)
            conn.commit()

        pii_warnings = detect_pii(headers, rows)

        return {
            "status": "success",
            "table_name": table_clean,
            "row_count": inserted_count,
            "columns": headers,
            "detected_types": col_types,
            "encoding": used_encoding,
            "pii_warnings": pii_warnings,
        }
    except Exception as e:
        return {
            "status": "error",
            "message": f"Failed to import file: {str(e)}"
        }


# รักษาความเข้ากันได้ย้อนหลัง 100% สำหรับโค้ดที่ import upload_csv_to_db
upload_csv_to_db = upload_file_to_db

